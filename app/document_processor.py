"""Document loading (PDF / DOCX / plain text, local or remote) and chunking."""

import re
import tempfile
from urllib.parse import urlparse

import requests


def get_file_extension_from_url(url: str) -> str:
    parsed = urlparse(url)
    path = parsed.path.lower()
    for ext in (".pdf", ".docx", ".doc", ".txt"):
        if path.endswith(ext):
            return ".docx" if ext == ".doc" else ext
    return ".pdf"  # default assumption


def download_to_tempfile(url: str) -> str:
    response = requests.get(url, timeout=30)
    if response.status_code != 200:
        raise ValueError(f"Failed to download document. Status: {response.status_code}")
    suffix = get_file_extension_from_url(url)
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        tmp.write(response.content)
        return tmp.name


def extract_text_from_pdf(path: str) -> str:
    import fitz  # PyMuPDF

    doc = fitz.open(path)
    try:
        return "\n".join(page.get_text() for page in doc)
    finally:
        doc.close()


def extract_text_from_docx(path: str) -> str:
    import docx

    doc = docx.Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text_from_document(path: str) -> str:
    ext = path.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        return extract_text_from_pdf(path)
    if ext in ("docx", "doc"):
        return extract_text_from_docx(path)
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def chunk_text(text: str, chunk_size: int = 500, overlap: int = 80) -> list[str]:
    """Character-based sliding-window chunking with real overlap between
    consecutive chunks, so context near a chunk boundary isn't lost."""
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return []
    if overlap >= chunk_size:
        overlap = chunk_size // 4

    chunks = []
    start = 0
    step = chunk_size - overlap
    while start < len(text):
        chunk = text[start:start + chunk_size].strip()
        if chunk:
            chunks.append(chunk)
        start += step
    return chunks
